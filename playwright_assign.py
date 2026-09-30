import json
import random
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright


BASE_DIR = Path(__file__).resolve().parent
CONTACTS_FILE = BASE_DIR / "contacts.xlsx"
REPORT_DATE = datetime.now().strftime("%Y-%m-%d")
REPORT_JSON = BASE_DIR / f"whatsapp_report_{REPORT_DATE}.json"
REPORT_EXCEL = BASE_DIR / f"whatsapp_report_{REPORT_DATE}.xlsx"
SCREENSHOTS_DIR = BASE_DIR / "screenshots"
PROFILE_DIR = BASE_DIR / "whatsapp_profile"
DEFAULT_MESSAGE = "Hello {name}"
LOGIN_TIMEOUT_MS = 5 * 60 * 1000


def random_delay(page):
    """Pause 2–5 seconds between browser actions."""
    page.wait_for_timeout(random.randint(2000, 5000))


def excel_cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_phone(phone):
    """Return an international phone number as digits or raise ValueError."""
    raw = excel_cell_text(phone)
    if not re.fullmatch(r"\+?[\d\s().-]+", raw):
        raise ValueError(f"Invalid phone number format: {raw!r}")

    digits = re.sub(r"\D", "", raw)
    if not 8 <= len(digits) <= 15 or digits.startswith("0"):
        raise ValueError(
            f"Phone must include a valid country code (8–15 digits): {raw!r}"
        )
    return digits


def read_contacts(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        headers = {
            str(cell.value).strip().lower(): index
            for index, cell in enumerate(sheet[1])
            if cell.value is not None
        }

        if "name" not in headers or "phone" not in headers:
            raise ValueError(
                "contacts.xlsx must contain Name and Phone columns."
            )

        contacts = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            name = values[headers["name"]] if headers["name"] < len(values) else None
            phone = values[headers["phone"]] if headers["phone"] < len(values) else None
            message = (
                values[headers["message"]]
                if "message" in headers and headers["message"] < len(values)
                else None
            )

            if name is None and phone is None:
                continue

            contacts.append({
                "name": excel_cell_text(name),
                "phone": excel_cell_text(phone),
                "message": excel_cell_text(message),
            })

        return contacts
    finally:
        workbook.close()


def save_reports(results, report_date):
    report_json = BASE_DIR / f"whatsapp_report_{report_date}.json"
    report_excel = BASE_DIR / f"whatsapp_report_{report_date}.xlsx"

    report_data = {
        "date": report_date,
        "results": results,
    }
    report_json.write_text(
        json.dumps(report_data, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append([
        "Name",
        "Phone",
        "Message",
        "Status",
        "Screenshot",
        "Last 3 Messages",
        "Error",
    ])

    for item in results:
        sheet.append([
            item["name"],
            item["phone"],
            item["message"],
            item["status"],
            item["screenshot"],
            "\n---\n".join(item["last_3_messages"]),
            item["error"],
        ])

    workbook.save(report_excel)


def open_contact(page, phone):
    """Look up and open the WhatsApp chat for an exact international number."""
    digits = normalize_phone(phone)
    page.goto(
        f"https://web.whatsapp.com/send?phone={digits}",
        wait_until="domcontentloaded",
    )

    try:
        page.wait_for_selector(
            '[data-testid="conversation-compose-box-input"], '
            '[contenteditable="true"][data-tab="10"], '
            '[contenteditable="true"][aria-label="Type a message"]',
            state="visible",
            timeout=20000,
        )
    except PlaywrightTimeoutError as error:
        body_text = page.locator("body").inner_text()
        if re.search(
            r"phone number.*(invalid|not on whatsapp)|not a valid phone",
            body_text,
            re.IGNORECASE,
        ):
            raise LookupError(
                "WhatsApp rejected the number or it is not registered."
            ) from error
        raise LookupError(
            "Could not open a WhatsApp chat for this number."
        ) from error

    random_delay(page)
    return digits


def message_box_for(page):
    selectors = (
        '[data-testid="conversation-compose-box-input"]',
        '[contenteditable="true"][data-tab="10"]',
        '[contenteditable="true"][aria-label="Type a message"]',
    )

    for selector in selectors:
        try:
            page.wait_for_selector(selector, state="visible", timeout=5000)
            return page.locator(selector).last
        except PlaywrightTimeoutError:
            continue

    raise RuntimeError('Could not locate the "Type a message" box.')


def wait_for_outgoing_message(page, previous_count):
    """Confirm a new outgoing bubble appears in the current chat."""
    page.wait_for_function(
        "previous => document.querySelectorAll('div.message-out').length > previous",
        arg=previous_count,
        timeout=30000,
    )


def extract_last_incoming_messages(page):
    incoming = page.locator("div.message-in")
    count = incoming.count()
    messages = []

    for index in range(max(0, count - 3), count):
        bubble = incoming.nth(index)
        text_parts = bubble.locator(
            '[data-testid="selectable-text"], .selectable-text'
        ).all_inner_texts()
        text = "\n".join(part.strip() for part in text_parts if part.strip())

        if not text:
            text = bubble.inner_text().strip()

        if text:
            messages.append(text)

    return messages[-3:]


def screenshot_sent_message(page, name, phone):
    SCREENSHOTS_DIR.mkdir(parents=True, exist_ok=True)
    label = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        name or phone or "contact",
    ).strip("_") or "contact"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    path = SCREENSHOTS_DIR / f"{label}_{timestamp}.png"

    outgoing = page.locator("div.message-out").last
    outgoing.wait_for(state="visible", timeout=10000)
    try:
        outgoing.screenshot(path=str(path))
    except Exception:
        page.screenshot(path=str(path), full_page=True)

    return str(path.relative_to(BASE_DIR))


def send_to_contact(page, contact):
    name = contact["name"]
    phone = contact["phone"]
    template = contact["message"] or DEFAULT_MESSAGE
    message = template.replace("{name}", name or phone)

    result = {
        "name": name,
        "phone": phone,
        "message": message,
        "status": "failed",
        "screenshot": "",
        "last_3_messages": [],
        "error": "",
    }

    try:
        if not phone:
            raise ValueError("Phone number is missing.")

        phone_digits = open_contact(page, phone)
        composer = message_box_for(page)
        outgoing_before = page.locator("div.message-out").count()

        composer.fill(message)
        random_delay(page)
        composer.press("Enter")
        wait_for_outgoing_message(page, outgoing_before)

        # A new outgoing bubble appeared, so sending was confirmed.
        result["status"] = "sent"
        random_delay(page)

        try:
            result["screenshot"] = screenshot_sent_message(
                page, name, phone_digits
            )
        except Exception as error:
            result["error"] = f"Screenshot failed: {error}"

        try:
            result["last_3_messages"] = extract_last_incoming_messages(page)
        except Exception as error:
            extraction_error = f"Message extraction failed: {error}"
            result["error"] = (
                f"{result['error']}; {extraction_error}"
                if result["error"]
                else extraction_error
            )

    except Exception as error:
        result["error"] = str(error)
        if result["status"] != "sent":
            result["status"] = "failed"

    return result


def main():
    report_date = datetime.now().strftime("%Y-%m-%d")

    if not CONTACTS_FILE.exists():
        print(f"Input file not found: {CONTACTS_FILE}", file=sys.stderr)
        return 1

    try:
        contacts = read_contacts(CONTACTS_FILE)
    except Exception as error:
        print(f"Could not read contacts: {error}", file=sys.stderr)
        return 1

    if not contacts:
        print("No contacts found in contacts.xlsx.", file=sys.stderr)
        return 1

    results = []
    context = None

    try:
        with sync_playwright() as playwright:
            context = playwright.chromium.launch_persistent_context(
                user_data_dir=str(PROFILE_DIR),
                channel="chrome",
                headless=False,
            )
            page = context.pages[0] if context.pages else context.new_page()
            page.goto("https://web.whatsapp.com", wait_until="domcontentloaded")

            print(
                "If this is the first run, scan the WhatsApp Web QR code "
                "in the browser."
            )

            page.wait_for_selector(
                "#pane-side",
                state="visible",
                timeout=LOGIN_TIMEOUT_MS,
            )

            for index, contact in enumerate(contacts, start=1):
                label = contact["name"] or contact["phone"] or "unknown contact"
                print(f"[{index}/{len(contacts)}] Processing {label}...")
                results.append(send_to_contact(page, contact))
                save_reports(results, report_date)
                random_delay(page)

    except PlaywrightTimeoutError:
        print(
            "WhatsApp login was not detected before timeout.",
            file=sys.stderr,
        )
        return 1
    except Exception as error:
        print(f"Automation stopped: {error}", file=sys.stderr)
        return 1
    finally:
        try:
            if context is not None:
                context.close()
        finally:
            save_reports(results, report_date)

    print(
        f"Reports saved to whatsapp_report_{report_date}.json and "
        f"whatsapp_report_{report_date}.xlsx."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
